"""Append the Milestone-2 partial-results slides to a copy of the project deck.

    python scripts/build_m2_slides.py   ->  ../Project_Presentation_whisenaj_M2.pptx (original untouched)

Every number on these slides comes from results/processed/*.csv|json and results/raw/*.json
(test split, seed 42 unless marked "3 seeds").
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_helpers import GARNET, INK, INK2, card, new_slide, picture, table, text  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
ROOT = REPO.parent
FIG = REPO / "results" / "figures"
SRC = ROOT / "Project_Presentation_whisenaj.pptx"
OUT = ROOT / "Project_Presentation_whisenaj_M2.pptx"
GREEN, AMBER, RED = RGBColor(0x00, 0x83, 0x00), RGBColor(0x9A, 0x5B, 0x00), RGBColor(0xB0, 0x1E, 0x2B)
B = {"bold": True}


def main():
    prs = Presentation(str(SRC))

    # 5 ---------------------------------------------------------------- hypotheses and verdicts
    s = new_slide(prs, "Partial Results: Hypotheses vs. What We Measured", (
        "This milestone asks whether adaptive resolution can cut the cost of RGB+IR early fusion without losing "
        "accuracy. Before any numbers, here is the scorecard. H4 is supported. H1 and H2 are partly or not "
        "supported, and I will explain why. H3, the learned controller, is next milestone's work, but this "
        "milestone measured the gap it has to close: the per-image oracle is 3.5 AP above the fixed-resolution "
        "frontier. Setup: M3FD test split (419 frames, sequence-grouped), multi-scale YOLO26n, seed 42; "
        "headline accuracy also over 3 seeds. Platforms: RTX 6000 Ada, and 4 pinned Xeon cores as the "
        "compute-bound stand-in until the Orin is up."))
    rows = [["", "Hypothesis (proposal)", "Result", "Verdict"],
            ["H1", "640→480 px costs < 2 mAP pts and cuts p95 > 30%",
             "p95 −32% on CPU, but −3.0 mAP pts (3 seeds)", "Partly"],
            ["H2", "Scene heuristic: ≥ 20% lower p95 at ≤ 2 pts; beats fixed-480 and random",
             "p95 not reduced (31.9 vs 31.1 ms); +0.3 AP over cost-matched random", "Not supported"],
            ["H3", "Learned controller closes ≥ 50% of the oracle gap",
             "Oracle gap measured: +3.5 AP over the fixed frontier at equal CPU cost", "Next milestone"],
            ["H4", "Savings larger on compute-bound hardware",
             "GPU latency flat (3.9–4.1 ms p50); CPU 12.6 → 26.6 ms (2.1×)", "Supported"]]
    t = table(s, 0.55, 1.05, 12.2, rows, col_w=[0.6, 4.6, 5.2, 1.8], size=14, row_h=0.72)
    for i, c in ((1, AMBER), (2, RED), (3, INK2), (4, GREEN)):
        run = t.cell(i, 3).text_frame.paragraphs[0].runs[0]
        run.font.color.rgb, run.font.bold = c, True
    text(s, 0.55, 5.0, 12.2, 1.6, [
        [("Takeaway: ", B), ("per-image headroom is real (oracle +3.5 AP), but scene-level controllers capture "
                             "almost none of it. On a desktop GPU, resolution is not a latency knob at all.", {})]],
         size=18)

    # 6 ---------------------------------------------------------------- what was built
    s = new_slide(prs, "What Was Built", (
        "Everything in the proposal's pipeline exists and runs end to end. The dashed boxes are reused open "
        "source; the rest is new. Ten detectors were trained: RGB-only, IR-only and 4-channel fusion at 640, "
        "fusion specialists at 480 and 320, and one multi-scale fusion model; the two headline models also on "
        "seeds 123 and 456. One GPU inference path is shared by evaluation and benchmarking, so the accuracy and "
        "latency numbers come from the same code. Three unit tests guard the risky parts: the 4-channel stem "
        "reproduces the COCO model exactly when IR is zero; our preprocessing matches Ultralytics' boxes; and our "
        "per-image AP matches pycocotools."))
    picture(s, ROOT / "figures" / "solution_dynamic_resolution_pipeline.png", 0.5, 0.9, 12.3, 3.3)
    for k, (big, label) in enumerate([("10", "detectors trained (6 configs, 3 seeds for headline)"),
                                      ("13", "resolution policies benchmarked on GPU + pinned CPU"),
                                      ("3", "correctness tests: stem, preprocessing, per-image AP"),
                                      ("0.8 → 0.1 ms", "Scene Analyzer cost on GPU after cost-aware trimming")]):
        card(s, 0.55 + k * 3.1, 4.55, 2.9, 1.75, big, label, big_size=30)

    # 7 ---------------------------------------------------------------- rigor fixes
    s = new_slide(prs, "Rigor: Three Things That Would Have Skewed the Results", (
        "Three problems turned up during setup; each would have biased the results. First, M3FD frames come from "
        "video sequences. Under the proposal's random split, 48% of test frames have a near-duplicate in train. "
        "Grouping frames into 767 sequences and splitting whole groups brings that to 0%. It also explains why "
        "published M3FD numbers run higher than ours. Second, my evaluator assumed every frame is 1024×768, but "
        "274 of 4,200 are smaller. Cross-checking against Ultralytics' own evaluator exposed a 5.8-point error; "
        "the remaining 1.5-point difference is fully explained by metric definition and resize kernel. Third, "
        "41% of all boxes are smaller than 32×32 pixels. Those are exactly the objects that downscaling destroys, "
        "so small-object AP and pedestrian recall are always reported separately."))
    card(s, 0.55, 1.0, 3.9, 1.7, "48% → 0%", "test frames with a near-duplicate in train (random vs. grouped split)")
    card(s, 0.55, 2.9, 3.9, 1.7, "5.8 pts", "evaluator error caught by cross-check (6.5% of frames not 1024×768)")
    card(s, 0.55, 4.8, 3.9, 1.7, "40.9%", "of 34,407 boxes are < 32² px: the objects downscaling kills")
    picture(s, REPO / "splits" / "group_contact_sheet.jpg", 4.8, 1.0, 8.0, 5.3)
    text(s, 4.8, 6.35, 8.0, 0.4, ["Rows = frames from one sequence group; whole groups go to one split."],
         size=11, color=INK2)

    # 8 ---------------------------------------------------------------- positive: fusion and multi-scale
    s = new_slide(prs, "Positive: Fusion Works, and One Model Serves All Resolutions", (
        "Two results support the design. Fusion helps: at 640 px, 4-channel early fusion beats RGB-only by 2.6 AP "
        "and IR-only by 7.3 AP. Pedestrian recall rises from 0.54 to 0.68, because IR sees people that RGB misses "
        "at night and in smoke. One multi-scale weight set is nearly free: over three seeds it gives up 0.7 AP at "
        "640 against the 640 specialist, and it is 2.2 AP better at 320. That is what makes resolution a runtime "
        "knob rather than a model swap, and it retires the proposal's risk that multi-scale training would cost "
        "more than 2 points."))
    rows = [["Test mAP@50-95 (3 seeds)", "320 px", "480 px", "640 px"],
            ["Fusion trained at 640 (B2)", "0.324 ± 0.003", "0.399 ± 0.005", "0.429 ± 0.002"],
            ["Fusion multi-scale (MS)", "0.346 ± 0.004", "0.394 ± 0.005", "0.422 ± 0.006"],
            ["RGB only (B0, seed 42)", "–", "–", "0.401"],
            ["IR only (B1, seed 42)", "–", "–", "0.354"]]
    table(s, 0.55, 1.1, 7.6, rows, col_w=[3.1, 1.5, 1.5, 1.5], size=14, row_h=0.62, bold_rows=(2,))
    card(s, 8.6, 1.1, 4.2, 1.6, "+2.6 / +7.3", "AP from fusion vs. RGB-only / IR-only at 640 px")
    card(s, 8.6, 2.9, 4.2, 1.6, "0.54 → 0.68", "People recall, RGB-only → fusion (IoU 0.5, conf 0.25)")
    card(s, 8.6, 4.7, 4.2, 1.6, "−0.7 / +2.2", "multi-scale vs. 640-specialist AP at 640 / at 320 px")
    text(s, 0.55, 4.5, 7.6, 1.9, [
        "Multi-scale training (one weight set) costs 0.7 AP at 640 px and gains 2.2 AP at 320 px.",
        "Proposal risk (\"multi-scale costs > 2 pts\") did not materialise."], size=15, bullets=True)

    # 9 ---------------------------------------------------------------- positive: headroom
    s = new_slide(prs, "Positive: The Per-Image Headroom Is Real", (
        "The per-image oracle picks, after the fact, the smallest resolution that stays within 0.05 AP of each "
        "image's best. Measured online on the pinned CPU, it runs 30% faster than fixed-640 on average and is "
        "0.8 AP more accurate. It can be more accurate because 320 px beats 640 px on 20% of test images, the same "
        "effect AdaScale reported for video. On the GPU, the left panel, every policy except v1 sits within about 0.4 ms: "
        "there is nothing to save there. Specialists are plotted at the multi-scale model's latency because they "
        "have identical compute."))
    picture(s, FIG / "F1_accuracy_latency.png", 0.4, 0.85, 9.0, 3.95)
    card(s, 9.7, 0.95, 3.1, 1.55, "−30%", "mean CPU latency, oracle vs. fixed-640 (18.6 vs 26.6 ms)")
    card(s, 9.7, 2.65, 3.1, 1.55, "+0.8 AP", "oracle (0.423) vs. fixed-640 (0.415)")
    card(s, 9.7, 4.35, 3.1, 1.55, "20%", "of test frames score higher at 320 than at 640 px")
    text(s, 0.55, 5.05, 9.0, 1.4, [
        [("Upper bound, not a deployable policy: ", B),
         ("the oracle is chosen with ground truth. It shows how much a perfect controller could save, which "
          "is the target for the controllers on the next slides.", {})]], size=14)

    # 10 --------------------------------------------------------------- safety
    s = new_slide(prs, "Safety: Lower Resolution Costs Small Objects and Pedestrians First", (
        "This is the safety figure from the proposal. Going from 640 to 320 px cuts small-object AP by a factor "
        "of three, from 0.19 to 0.06, and pedestrian recall from 0.69 to 0.56. Aggregate mAP falls only 17% over "
        "the same range, so aggregate mAP hides the damage. The examples on the right are real test frames: at "
        "320 px the detector finds 2 of 9 small pedestrians in the first frame; at 640 px it finds 7. Any "
        "controller must be judged on these metrics, not only on mAP."))
    picture(s, FIG / "F4_small_object_safety.png", 0.4, 0.9, 6.6, 3.3)
    picture(s, FIG / "F5_detection_examples.png", 7.1, 0.85, 5.9, 5.9)
    card(s, 0.55, 4.5, 3.1, 1.75, "3.1×", "small-object AP loss, 640 → 320 px (0.191 → 0.062)")
    card(s, 3.85, 4.5, 3.1, 1.75, "0.69 → 0.56", "People recall, 640 → 320 px")

    # 11 --------------------------------------------------------------- negative: controllers
    s = new_slide(prs, "Negative: Scene-Level Controllers Don't Find the Headroom", (
        "This is the main negative result. The first controller, a decision tree trained to predict the oracle's "
        "resolution, is no better than always guessing the most common answer (41% accuracy either way). "
        "Diagnosis: no scene statistic correlates with how much an image gains from resolution; the largest "
        "correlation is 0.16. The gain depends on small objects that a thumbnail cannot see. The detector's own "
        "320-px output can see them (correlation 0.39), but running that probe costs 44% of a 640 pass on CPU. "
        "Every cascade I tried loses, including cheaper 192 and 256 px probes. The best deployable controller is "
        "one rule, fit on train and chosen on val: mean luminance ≤ 142 → 480 px, else 640 px. It beats a "
        "cost-matched random mix by only 0.3 AP, and online it does not shorten the latency tail. All variants were selected on val and scored once on test, and "
        "all of them appear in this figure, including the failures."))
    picture(s, FIG / "F6_controller_design_space.png", 0.4, 0.9, 7.6, 4.65)
    text(s, 8.25, 0.95, 4.7, 5.6, [
        [("v1 classification tree: ", B), ("CV accuracy 0.41 = majority baseline", {})],
        [("Why: ", B), ("scene features vs. per-image gain |ρ| ≤ 0.16; true small-object share ρ = 0.37", {})],
        [("Probe cascade: ", B), ("320-px detections ρ = 0.39, but the probe costs 44% of a 640 pass → loses", {})],
        [("v2 luminance rule: ", B), ("0.405 vs 0.403 ± 0.002 for cost-matched random (+0.3 AP)", {})],
        [("Online (CPU): ", B), ("v2 p95 31.9 ms vs 31.1 for fixed-640; the 640 frames set the tail", {})],
        [("Captured headroom: ", B), ("≈ +0.3 of the oracle's +3.5 AP (all 12 sweep points just above the random frontier)", {})]],
         size=14, bullets=True, space_after=9)

    # 12 --------------------------------------------------------------- negative: GPU flat
    s = new_slide(prs, "Negative (and H4): On a Desktop GPU, Resolution Isn't a Latency Knob", (
        "The second negative result is also the evidence for H4. On the RTX 6000 Ada at batch size 1, YOLO26n "
        "takes about 4 ms at every resolution. The GPU spends its time on per-kernel launch overhead, not "
        "arithmetic, so shrinking the input saves nothing, and any controller only adds overhead. Energy does "
        "still fall, by 23% at 320 px. On 4 pinned CPU cores the same model is compute-bound and latency scales "
        "2.1× from 320 to 640. That is the regime an edge device lives in, which is why the Orin measurements "
        "next milestone matter. The analyzer was also trimmed to the one feature the rule uses: from 0.81 to "
        "0.11 ms on GPU, and from 6.9 to 2.3 ms on CPU."))
    picture(s, FIG / "F3_stage_latency.png", 0.4, 0.9, 6.3, 3.4)
    picture(s, FIG / "F3_stage_latency_cpu4.png", 6.7, 0.9, 6.3, 3.4)
    card(s, 0.55, 4.6, 3.0, 1.7, "3.9–4.1 ms", "GPU p50 at 320 / 480 / 640 px: flat (launch-bound)")
    card(s, 3.75, 4.6, 3.0, 1.7, "−23%", "GPU energy per frame at 320 vs. 640 px (0.31 vs 0.40 J)")
    card(s, 6.95, 4.6, 3.0, 1.7, "2.1×", "CPU latency 320 → 640 px (12.6 → 26.6 ms): compute-bound")
    card(s, 10.15, 4.6, 2.65, 1.7, "6.9 → 2.3 ms", "analyzer on CPU after cost-aware trimming")

    # 13 --------------------------------------------------------------- related work
    s = new_slide(prs, "How These Results Relate to Prior Work", (
        "Our findings line up with prior work, and the differences explain our controller results. AdaScale "
        "reported that downscaling sometimes improves accuracy; we see the same on 20% of frames. But AdaScale "
        "works on video, where the previous frame's detections are a free probe. On single frames, the probe "
        "costs a full forward pass, and that is exactly where our cascade failed. DRNet reports compute savings "
        "in FLOPs on classification. Our GPU result shows why FLOPs mislead: fewer FLOPs gave no latency change. "
        "RTScale and Chanakya treat resolution as a runtime decision under a deadline; our CPU/GPU contrast "
        "supports their point that the right policy depends on the hardware's cost model."))
    rows = [["Work", "What they found", "What we see on M3FD RGB+IR"],
            ["AdaScale (MLSys '19)", "Per-frame scale from the previous frame; downscaling can raise mAP",
             "Confirmed: 320 beats 640 on 20% of frames. Without a free temporal probe, predicting which frames fails"],
            ["DRNet (NeurIPS '21)", "−34% FLOPs at equal accuracy (classification)",
             "FLOPs ≠ latency: GPU time flat across 4× fewer pixels"],
            ["RTScale (ECRTS '22), Chanakya (NeurIPS '23)", "Runtime resolution under deadlines; tail latency matters",
             "Mixed policies keep the 640-px tail; policy must be tuned per platform (H4)"],
            ["TarDAL / M3FD (CVPR '22)", "Multi-scenario RGB+IR benchmark",
             "Sequence leakage: 48% of test frames near-duplicated under random split"]]
    table(s, 0.55, 1.05, 12.2, rows, col_w=[2.9, 4.3, 5.0], size=13, row_h=0.95)

    # 14 --------------------------------------------------------------- next steps
    s = new_slide(prs, "Next Steps", (
        "Each next step follows from a result on the earlier slides. First, bring up the Orin and run the same "
        "benchmark session there; the CPU results predict that resolution matters on the edge. Second, the "
        "learned controller of RQ3: a small CNN on the four-channel thumbnail. Its training labels must come "
        "from models that did not see those images, because train-split labels are inflated. Third, find a probe "
        "that costs nothing: the previous frame's detections in video, or features the backbone has already "
        "computed. Fourth, extend three seeds to every configuration and add energy per correct detection. "
        "Scope decisions made so far: the proposal's random split was replaced by the grouped split, and the "
        "Orin is replaced by a pinned CPU until it is up."))
    steps = [("1", "Jetson Orin bring-up", "Record JetPack and power mode; rerun bench_session.sh unchanged. "
                                           "The CPU results predict resolution does matter at the edge."),
             ("2", "Learned controller (RQ3)", "Small CNN on the 4-channel thumbnail. Train it on labels from "
                                               "models that never saw those images (train-split labels are inflated)."),
             ("3", "A free probe", "Previous-frame detections (video) or already-computed backbone features. "
                                   "The probe signal exists (ρ = 0.39); its cost is the problem."),
             ("4", "Complete the study", "3 seeds for all configs; energy per correct detection; "
                                         "A2 feature ablation; final F1–F5 on both platforms.")]
    for k, (n, head, body) in enumerate(steps):
        y = 1.0 + k * 1.42
        card(s, 0.55, y, 1.1, 1.2, n, "", big_size=32)
        text(s, 1.9, y + 0.05, 10.9, 0.45, [head], size=18, bold=True, color=GARNET)
        text(s, 1.9, y + 0.5, 10.9, 0.8, [body], size=14, color=INK)

    prs.save(str(OUT))
    print("wrote", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
