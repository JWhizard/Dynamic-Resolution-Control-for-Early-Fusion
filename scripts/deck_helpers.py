"""python-pptx helpers for appending Milestone-2 slides in the style of the existing deck
(layout "iMSEL Slide": garnet Arial title band, white body, Arial text)."""

from __future__ import annotations

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

GARNET = RGBColor(0x73, 0x00, 0x0A)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
TINT = RGBColor(0xF2, 0xF0, 0xF0)  # theme light-gray tint used for callout cards
FONT = "Arial"
SLIDE_W, SLIDE_H = 13.333, 7.5
TOP = 0.85  # below the title band


def layout(prs, name="iMSEL Slide"):
    return next(l for l in prs.slide_layouts if l.name == name)


def new_slide(prs, title: str, notes: str):
    s = prs.slides.add_slide(layout(prs))
    s.shapes.title.text = title
    for ph in s.placeholders:
        if ph.placeholder_format.type is not None and ph.placeholder_format.idx not in (0,) and \
                "Number" not in ph.name and ph.has_text_frame and not ph.text_frame.text:
            ph._element.getparent().remove(ph._element)  # drop empty body placeholders
    s.notes_slide.notes_text_frame.text = notes
    return s


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bullets=False, space_after=6):
    """paras: list of str or list of (str, dict) runs; '**x**' style not parsed, use tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for k, p in enumerate(paras):
        para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [(p, {})]
        if bullets:
            _bullet(para)
        for t, st in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.italic = st.get("italic", False)
            f.color.rgb = st.get("color", color)
    return tb


def _bullet(para):
    from pptx.oxml.ns import qn
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.25))))
    pPr.set("indent", str(-Emu(Inches(0.22))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def picture(slide, path, x, y, w=None, h=None):
    """Fit inside the (w,h) box preserving aspect; centred in the box."""
    iw, ih = Image.open(path).size
    if w and h:
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                        Inches(pw), Inches(ph))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None,
                                    Inches(h) if h else None)


def card(slide, x, y, w, h, big: str, label: str, big_color=GARNET, big_size=34):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = TINT
    sh.line.fill.background()
    sh.shadow.inherit = False
    size = fit_size(big, w - 0.45, big_size)
    text(slide, x + 0.15, y + 0.08, w - 0.3, h * 0.50, [big], size=size, color=big_color, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    if label:
        text(slide, x + 0.15, y + h * 0.60, w - 0.3, h * 0.38, [label], size=12, color=INK2)
    return sh


_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"  # metric-compatible with Arial Bold


def fit_size(s: str, width_in: float, max_pt: int, min_pt: int = 16) -> int:
    """Largest point size <= max_pt at which `s` fits on one line of `width_in` inches."""
    from PIL import ImageFont
    for pt in range(max_pt, min_pt - 1, -1):
        if ImageFont.truetype(_BOLD, pt).getlength(s) / 72 <= width_in:
            return pt
    return min_pt


def table(slide, x, y, w, rows, col_w=None, size=12, header_fill=GARNET, row_h=0.36, bold_rows=()):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    if col_w:
        for k, cw in enumerate(col_w):
            gt.columns[k].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = i == 0 or i in bold_rows
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (TINT if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
    return gt
